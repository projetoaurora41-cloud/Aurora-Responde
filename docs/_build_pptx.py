# -*- coding: utf-8 -*-
"""Monta um PPTX do deck v2 com os 3 slides novos (RAG x2 + Encaminhamentos)
NATIVOS/editáveis (caixas de texto, cores gov.br) e os demais como imagem.
"""
import pymupdf
from pptx import Presentation
from pptx.util import Mm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import tempfile, os

FOLDER = r"C:\Users\herna\Downloads\Nova pasta (5)"
ORIG   = FOLDER + r"\Apresentacao_Aurora_MP_2026_09_18.pdf"
OUT    = FOLDER + r"\Apresentacao_Aurora_MP_v2_editavel.pptx"

# ---- paleta gov.br ----
BLUE=RGBColor(0x13,0x51,0xB4); DARK=RGBColor(0x07,0x1D,0x41); BLUE6=RGBColor(0x0c,0x32,0x6f)
YEL=RGBColor(0xFF,0xCD,0x07); GREEN=RGBColor(0x16,0x88,0x21); RED=RGBColor(0xE5,0x22,0x07)
MUTED=RGBColor(0x5b,0x6b,0x7b); TEXT=RGBColor(0x1b,0x2a,0x3a); BORDER=RGBColor(0xcf,0xda,0xea)
LBLUE=RGBColor(0xee,0xf3,0xfb); CARD=RGBColor(0xff,0xff,0xff); PANEL=RGBColor(0xf3,0xf7,0xfd)
YELBG=RGBColor(0xff,0xf7,0xe0); DARKTX=RGBColor(0x33,0x27,0x0a)
FONT="Calibri"

prs = Presentation(); prs.slide_width=Mm(297); prs.slide_height=Mm(210)
BLANK = prs.slide_layouts[6]

def slide():
    return prs.slides.add_slide(BLANK)

def rect(s,x,y,w,h,fill=None,line=None,lw=0.75,rounded=False):
    shp=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                           Mm(x),Mm(y),Mm(w),Mm(h))
    shp.shadow.inherit=False
    if fill is None: shp.fill.background()
    else: shp.fill.solid(); shp.fill.fore_color.rgb=fill
    if line is None: shp.line.fill.background()
    else: shp.line.color.rgb=line; shp.line.width=Pt(lw)
    if rounded:
        try: shp.adjustments[0]=0.10
        except Exception: pass
    return shp

def circle(s,x,y,d,fill):
    shp=s.shapes.add_shape(MSO_SHAPE.OVAL,Mm(x),Mm(y),Mm(d),Mm(d))
    shp.shadow.inherit=False; shp.fill.solid(); shp.fill.fore_color.rgb=fill
    shp.line.fill.background(); return shp

def tb(s,x,y,w,h,paras,anchor=MSO_ANCHOR.TOP):
    """paras: lista de paragrafos; cada paragrafo = dict(runs=[(txt,sz,bold,color,italic)],
       align, space_after)."""
    box=s.shapes.add_textbox(Mm(x),Mm(y),Mm(w),Mm(h)); tf=box.text_frame
    tf.word_wrap=True; tf.vertical_anchor=anchor
    tf.margin_left=Mm(1); tf.margin_right=Mm(1); tf.margin_top=Mm(0.5); tf.margin_bottom=Mm(0.5)
    for i,p in enumerate(paras):
        par=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        par.alignment=p.get("align",PP_ALIGN.LEFT)
        if "space_after" in p: par.space_after=Pt(p["space_after"])
        if "line" in p: par.line_spacing=p["line"]
        for (txt,sz,bold,color,ital) in p["runs"]:
            r=par.add_run(); r.text=txt; f=r.font
            f.size=Pt(sz); f.bold=bold; f.italic=ital; f.name=FONT; f.color.rgb=color
    return box

def R(txt,sz=10,bold=False,color=TEXT,ital=False): return (txt,sz,bold,color,ital)
def P(runs,align=PP_ALIGN.LEFT,space_after=2,line=1.0):
    if isinstance(runs,tuple): runs=[runs]
    return dict(runs=runs,align=align,space_after=space_after,line=line)

def header(s, gov=True):
    rect(s,14,11,9,9,fill=BLUE,rounded=True)
    tb(s,15.5,12.2,7,7,[P(R("A",13,True,RGBColor(255,255,255)),align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,25,10.5,150,10,[
        P(R("PROJETO AURORA",13,True,DARK),space_after=0),
        P(R("Análise Unificada de Registros para Orientar Respostas e Ações de Proteção à Criança e ao Adolescente",8,False,MUTED),space_after=0),
    ])
    tb(s,200,11,83,9,[
        P(R("MINISTÉRIO DA JUSTIÇA E SEGURANÇA PÚBLICA",8.5,True,DARK),align=PP_ALIGN.RIGHT,space_after=0),
        P(R("UFS · UFBA",8.5,False,MUTED),align=PP_ALIGN.RIGHT,space_after=0),
    ])
    rect(s,14,21.5,269,0.6,fill=BLUE)

def beta(s,x,y):
    b=rect(s,x,y,26,7,fill=YEL,rounded=True)
    tb(s,x,y+0.4,26,6.5,[P(R("BETA · v0.1.0",10,True,DARKTX),align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)

def numstrip(s, cells):
    rect(s,0,197,297,13,fill=BLUE)
    n=len(cells); cw=297/n
    for i,(v,l) in enumerate(cells):
        tb(s,i*cw,198.3,cw,10,[
            P(R(v,14,True,RGBColor(255,255,255)),align=PP_ALIGN.CENTER,space_after=0),
            P(R(l,8,False,RGBColor(0xe6,0xee,0xfb)),align=PP_ALIGN.CENTER,space_after=0),
        ])

# ============================================================ SLIDE A: RAG como funciona
def slide_rag():
    s=slide(); header(s)
    tb(s,14,26,210,12,[P([R("Como funciona o ",26,True,DARK),R("AURORA Responde",26,True,BLUE),R(" (RAG)",26,True,DARK)])])
    beta(s,228,29)
    tb(s,14,39,269,7,[P([R("O assistente ",11,False,MUTED),R("recupera a informação na fonte oficial",11,True,TEXT),
                         R(" e o modelo ",11,False,MUTED),R("responde citando essa fonte",11,True,TEXT),
                         R(" — verificável, sem “inventar”.",11,False,MUTED)])])
    # fluxo 5 passos
    steps=[("1","Pergunta","Em linguagem natural (até 2.000 caracteres)."),
           ("2","Guardrails","Mascara dados pessoais, bloqueia usos indevidos e, em risco, encaminha ao Conselho Tutelar / Disque 100."),
           ("3","Roteador (Qwen3)","Interpreta e escolhe a ferramenta da fonte certa (até 6 rodadas)."),
           ("4","Recupera + gera","Busca o trecho na fonte e redige a resposta ancorada, citando a origem."),
           ("5","Resposta","Prosa clara + a ferramenta consultada, a fonte e os tokens.")]
    x=14; y=47; w=51; gap=2.5
    for (n,h,t) in steps:
        rect(s,x,y,w,26,fill=LBLUE,line=RGBColor(0xd9,0xe4,0xf6),rounded=True)
        circle(s,x+2.5,y-2.5,6,YEL); tb(s,x+2.5,y-2.3,6,6,[P(R(n,11,True,DARKTX),align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
        tb(s,x+2.5,y+2.2,w-5,22,[P(R(h,11,True,BLUE6),space_after=1),P(R(t,8.5,False,RGBColor(0x33,0x46,0x5a),),line=1.02,space_after=0)])
        x+=w+gap
    # modelos (esquerda)
    tb(s,14,77,160,6,[P(R("MODELOS DE IA — LOCAIS, VIA OLLAMA",10,True,MUTED))])
    models=[("Qwen3:8b — Orquestrador/roteador","Interpreta a pergunta, escolhe a ferramenta e redige a resposta.",
             "8,2 bi parâmetros · janela 40.960 tokens","→ a janela é a memória: pergunta + dados recuperados + histórico."),
            ("Chronos-2 — Modelo de série temporal","Valida a evolução dos registros contra o histórico (retrospectiva).",
             "forecaster do comitê · fidelidade retrospectiva","→ indica se a tendência é real, e não ruído."),
            ("Jurema-7B — LLM jurídico","Redige o parecer ancorado no corpus de leis, citando lei e artigo.",
             "7,6 bi parâmetros · janela 32.768 tokens · até 450 tokens/resposta","→ o corpus corrige o que o modelo erraria sozinho.")]
    my=83.5; mh=36;
    for (nm,desc,specs,why) in models:
        rect(s,14,my,160,mh-2.5,fill=CARD,line=BORDER,rounded=True)
        tb(s,17,my+2,154,mh-5,[
            P(R(nm,11.5,True,DARK),space_after=1),
            P(R(desc,9,False,RGBColor(0x33,0x46,0x5a)),space_after=1,line=1.02),
            P(R(specs,9,True,BLUE6),space_after=1),
            P(R(why,9,False,GREEN),space_after=0,line=1.02),
        ])
        my+=mh
    # caixa RAG (direita, escura) + por que importa
    rect(s,178,83.5,105,40,fill=DARK,rounded=True)
    tb(s,181,85.5,99,37,[
        P(R("O que é RAG (Retrieval-Augmented Generation)",11.5,True,RGBColor(255,255,255)),space_after=2),
        P([R("Recupera",9.5,True,YEL),R(" o trecho da fonte e ",9.5,False,RGBColor(0xd6,0xe2,0xfb)),
           R("aumenta",9.5,True,YEL),R(" o modelo, que ",9.5,False,RGBColor(0xd6,0xe2,0xfb)),
           R("gera",9.5,True,YEL),R(" a resposta ",9.5,False,RGBColor(0xd6,0xe2,0xfb)),
           R("citando a origem",9.5,True,YEL),
           R(". Os números crus vão para as tabelas/gráficos; o modelo recebe só o resumo.",9.5,False,RGBColor(0xd6,0xe2,0xfb))],line=1.05,space_after=2),
        P(R("Fontes: SINAN · Séries temporais · ECA/legislação · SIPIA-CT · Documentação",8.5,True,RGBColor(255,255,255)),space_after=0),
    ])
    rect(s,178,125.5,105,45,fill=YELBG,line=RGBColor(0xf0,0xe2,0xb6),rounded=True)
    tb(s,181,127.5,99,42,[
        P(R("Por que essa arquitetura importa",11,True,RGBColor(0x8a,0x6d,0x00)),space_after=2),
        P([R("✓ Privacidade e soberania: ",9.5,True,RGBColor(0x5a,0x4a,0x00)),R("modelos locais; dados sensíveis de crianças não saem para serviços externos.",9.5,False,RGBColor(0x4a,0x3d,0x12))],line=1.05,space_after=2),
        P([R("✓ Respostas verificáveis: ",9.5,True,RGBColor(0x5a,0x4a,0x00)),R("ancoradas na fonte e com citação — menos alucinação.",9.5,False,RGBColor(0x4a,0x3d,0x12))],line=1.05,space_after=2),
        P([R("✓ Sem custo por consulta: ",9.5,True,RGBColor(0x5a,0x4a,0x00)),R("modelos abertos e próprios.",9.5,False,RGBColor(0x4a,0x3d,0x12))],line=1.05,space_after=0),
    ])
    numstrip(s,[("≤ 2.000","caracteres por pergunta"),("24 mensagens","memória da conversa"),
                ("6 rodadas","de ferramentas (temp. 0,2)"),("~4.100 + 250","tokens (pergunta + resposta)"),
                ("12 normas","no corpus jurídico")])

# ============================================================ SLIDE B: conteúdo do RAG
def slide_conteudo():
    s=slide(); header(s)
    tb(s,14,26,220,12,[P([R("O que está no ",26,True,DARK),R("RAG",26,True,BLUE),R(" — conteúdo consultado",26,True,DARK)])])
    beta(s,250,29)
    tb(s,14,39,269,7,[P([R("Cinco fontes oficiais, uma única interface. O modelo ",11,False,MUTED),
                         R("recupera o trecho na fonte",11,True,TEXT),R(" e responde ",11,False,MUTED),
                         R("citando a origem",11,True,TEXT),R(".",11,False,MUTED)])])
    fontes=[("SINAN — dados de saúde","~3,08 mi de notificações (2014–2024). Brasil/UF/município, ranking, contagens e perfis."),
            ("Séries temporais","Evolução no tempo; o Chronos-2 valida se a tendência é real."),
            ("Legislação (ECA e normas)","Corpus curado de leis; parecer do Jurema-7B citando lei e artigo."),
            ("SIPIA-CT","Registros dos Conselhos Tutelares por UF/ano, por dimensão."),
            ("Documentação AURORA","Como a plataforma funciona: metodologia e camada semântica.")]
    x=14;y=48;w=51;gap=2.5
    for (h,t) in fontes:
        rect(s,x,y,w,30,fill=CARD,line=BORDER,rounded=True)
        tb(s,x+2.2,y+2,w-4.4,26,[P(R(h,10.5,True,DARK),space_after=1,line=1.0),P(R(t,8.3,False,RGBColor(0x33,0x46,0x5a)),line=1.03,space_after=0)])
        x+=w+gap
    # painel corpus juridico
    rect(s,14,82,269,104,fill=PANEL,line=BORDER,rounded=True)
    tb(s,18,84.5,230,7,[P([R("Corpus jurídico",15,True,DARK),R("   — normas indexadas, geradas pelo Jurema-7B com citação de lei e artigo",10,False,MUTED)])])
    b=rect(s,250,84.8,29,7,fill=YELBG,line=RGBColor(0xf0,0xe2,0xb6),rounded=True)
    tb(s,250,85.1,29,6.5,[P(R("12 normas",10,True,RGBColor(0x8a,0x6d,0x00)),align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
    laws=[("CF/88 — art. 227 (+§4º)","proteção integral; punição do abuso sexual",BLUE),
          ("CP — art. 136","maus-tratos",RED),
          ("CP — art. 217-A","estupro de vulnerável (<14 anos)",RED),
          ("CP — arts. 218/218-A/218-B","corrupção e exploração sexual",RED),
          ("CP — art. 146-A","bullying e cyberbullying",RED),
          ("ECA — Lei 8.069/1990","proteção integral; comunicação ao CT",BLUE),
          ("Lei Henry Borel — 14.344/2022","violência doméstica e familiar",BLUE),
          ("Escuta Protegida — 13.431/2017","depoimento especial; sem revitimização",BLUE),
          ("Menino Bernardo — 13.010/2014","proíbe o castigo físico",GREEN),
          ("Lei 14.811/2024","escolas + crimes digitais; hediondez",RED),
          ("Lei 15.487/2026","violência sexual digital + IA (deepfakes)",RED),
          ("Denúncia / Boletim de Ocorrência","Conselho Tutelar · Disque 100 · delegacia (online)",GREEN)]
    cols=4; cw=64.5; ch=19; gx=2.2; gy=2.2; x0=18; y0=93.5
    for idx,(nm,desc,accent) in enumerate(laws):
        r=idx//cols; c=idx%cols
        cx=x0+c*(cw+gx); cy=y0+r*(ch+gy)
        rect(s,cx,cy,cw,ch,fill=CARD,line=RGBColor(0xd9,0xe4,0xf6),rounded=True)
        rect(s,cx,cy,1.2,ch,fill=accent)  # borda-esquerda colorida
        tb(s,cx+2.5,cy+1.5,cw-3.5,ch-3,[P(R(nm,9.8,True,DARK),space_after=0,line=1.0),P(R(desc,8,False,MUTED),space_after=0,line=1.0)])
    # SIPIA indicadores
    tb(s,18,170,265,10,[P([R("SIPIA-CT — indicadores:  ",10,True,DARK),
        R("Sexo · Cor · Faixa Etária · Direito Violado · Agente Violador",10,False,BLUE6),
        R("        ✓ resposta ancorada e verificável — números crus vão para os gráficos",9.5,True,GREEN)])])
    numstrip(s,[("5 fontes","numa única interface"),("~3,08 mi","notificações (SINAN)"),
                ("12 normas","no corpus jurídico"),("5 indicadores","do SIPIA-CT"),
                ("citação","da fonte em cada resposta")])

# ============================================================ SLIDE C: Encaminhamentos
def slide_enc():
    s=slide(); header(s)
    tb(s,14,27,260,12,[P(R("AURORA Responde na prática",27,True,DARK))])
    rect(s,14,44,269,150,fill=LBLUE,rounded=True)
    circle(s,20,49,9,BLUE); tb(s,20,49.2,9,9,[P(R("3",15,True,RGBColor(255,255,255)),align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,32,48.5,245,10,[P([R("ENCAMINHAMENTOS SITUACIONAIS   ",13,True,DARK),
                            R("“Meu vizinho maltrata uma criança e quero denunciar”",13,False,BLUE6)],align=PP_ALIGN.LEFT)])
    # card do chat
    rect(s,22,60,240,126,fill=CARD,rounded=True,line=RGBColor(0xe3,0xe9,0xf2))
    # usuario
    circle(s,27,64,7,BLUE6); tb(s,27,64.3,7,7,[P(R("V",11,True,RGBColor(255,255,255)),align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,37,63,200,6,[P([R("Você ",11,True,TEXT),R("· 11:34",10,False,MUTED)])])
    rect(s,37,69,150,8.5,fill=RGBColor(0xf2,0xf5,0xfa),line=RGBColor(0xe3,0xe9,0xf2),rounded=True)
    tb(s,40,69.6,145,7.5,[P(R("Meu vizinho maltrata uma criança e quero denunciar",11,False,RGBColor(0x33,0x46,0x5a)))],anchor=MSO_ANCHOR.MIDDLE)
    # aurora
    circle(s,27,82,7,GREEN); tb(s,27,82.3,7,7,[P(R("A",11,True,RGBColor(255,255,255)),align=PP_ALIGN.CENTER)],anchor=MSO_ANCHOR.MIDDLE)
    tb(s,37,81,200,6,[P([R("Aurora ",11,True,TEXT),R("· 11:34",10,False,MUTED)])])
    tb(s,37,87,222,96,[
        P([R("Pelos relatos envolvendo criança ou adolescente, o caminho adequado é procurar o ",11,False,RGBColor(0x2b,0x3a,0x4b)),
           R("Conselho Tutelar",11,True,DARK),R(" do município — órgão local responsável por aplicar medidas de proteção quando direitos são ameaçados ou violados.",11,False,RGBColor(0x2b,0x3a,0x4b))],line=1.25,space_after=6),
        P([R("Como maus-tratos, abuso e exploração são ",11,False,RGBColor(0x2b,0x3a,0x4b)),R("crimes",11,True,DARK),
           R(" (por exemplo, maus-tratos — art. 136 do Código Penal), você também pode ",11,False,RGBColor(0x2b,0x3a,0x4b)),
           R("registrar um boletim de ocorrência (B.O.)",11,True,DARK),
           R(": presencialmente na delegacia (de preferência a Delegacia de Proteção à Criança e ao Adolescente) ou pela ",11,False,RGBColor(0x2b,0x3a,0x4b)),
           R("Delegacia Eletrônica (online)",11,True,DARK),
           R(" do seu estado. O B.O. formaliza a denúncia e permite a investigação policial.",11,False,RGBColor(0x2b,0x3a,0x4b))],line=1.25,space_after=6),
        P([R("Você também pode denunciar de forma anônima no ",11,False,RGBColor(0x2b,0x3a,0x4b)),R("Disque 100",11,True,DARK),
           R(" (telefone 100), gratuito e 24 horas. Em emergência imediata, ligue ",11,False,RGBColor(0x2b,0x3a,0x4b)),R("190",11,True,DARK),R(".",11,False,RGBColor(0x2b,0x3a,0x4b))],line=1.25,space_after=6),
        P(R("O Aurora Responde trata de dados e legislação; não substitui o atendimento do Conselho Tutelar, da delegacia nem de outros serviços de proteção.",11,False,RGBColor(0x2b,0x3a,0x4b)),line=1.25,space_after=0),
    ])

# ---- montagem: imagens (originais) + nativos, na ordem do v2 ----
doc=pymupdf.open(ORIG); tmp=tempfile.mkdtemp(prefix="aur_")
def img_slide(page_index):
    pix=doc[page_index].get_pixmap(dpi=170); p=os.path.join(tmp,f"o{page_index}.png"); pix.save(p)
    s=slide(); s.shapes.add_picture(p,Mm(0),Mm(0),width=Mm(297),height=Mm(210))

for i in range(0,10): img_slide(i)     # originais 1..10
slide_rag()                            # 11
slide_conteudo()                       # 12
img_slide(10); img_slide(11)           # originais 11,12 -> 13,14 (na prática DADOS, LEGISLACAO)
slide_enc()                            # 15 (substitui encaminhamentos)
img_slide(13); img_slide(14)           # originais 14,15 -> 16,17 (outros exemplos, encerramento)

prs.save(OUT)
print("OK:",OUT,"| slides:",len(prs.slides._sldIdLst),"| tam MB:",round(os.path.getsize(OUT)/1e6,1))
